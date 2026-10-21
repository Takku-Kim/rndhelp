"""참고자료용 엑셀 내보내기 (AI 학습데이터)."""
from __future__ import annotations

import io

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core import loader


def train_frame() -> pd.DataFrame:
    d = loader.load_csv("train_ledger.csv").copy()
    d["account"] = d["account"].astype(str)
    nm = dict(zip(loader.accounts()["code"].astype(str), loader.accounts()["name"]))
    d.insert(1, "계정명", d["account"].map(nm))
    d.insert(0, "No", range(1, len(d) + 1))
    d.columns = ["No", "K-water 계정코드", "계정명", "적요", "거래처", "금액(원)", "RCMS 비목", "RCMS 세목"]
    return d


def train_xlsx() -> bytes:
    d = train_frame()
    sm = d.groupby(["RCMS 비목", "RCMS 세목"]).agg(건수=("No", "count"), 금액합계=("금액(원)", "sum")).reset_index()
    cm = d.groupby(["K-water 계정코드", "계정명", "RCMS 비목"]).size().reset_index(name="건수")
    info = pd.DataFrame({
        "항목": ["파일 용도", "구성", "데이터 성격", "컬럼 설명", "출처 파일"],
        "내용": ["RND-HELP 비목 분류 AI 학습데이터 참고자료",
               f"학습데이터 {len(d):,}건 / 비목·세목별 요약 / 계정×비목 대응",
               "가상(합성) 데이터 — 실제 K-water 집행내역이 아님. 실제 내역으로 교체 예정",
               "계정코드·적요·거래처 → RCMS 비목·세목 (정답 라벨은 룰 기반으로 부여)",
               "data/train_ledger.csv"]})
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        info.to_excel(w, sheet_name="안내", index=False)
        d.to_excel(w, sheet_name="학습데이터", index=False)
        sm.to_excel(w, sheet_name="비목세목 요약", index=False)
        cm.to_excel(w, sheet_name="계정×비목", index=False)
        thin = Side(style="thin", color="DDE3EA")
        for ws in w.book.worksheets:
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="0B3B6E")
                c.alignment = Alignment(horizontal="center", vertical="center")
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
                    c.alignment = Alignment(vertical="center", wrap_text=(ws.title == "안내"))
            for i, col in enumerate(ws.columns, 1):
                m = max(sum(2 if ord(ch) > 127 else 1 for ch in str(c.value or "")) for c in col)
                ws.column_dimensions[get_column_letter(i)].width = min(max(m + 3, 8), 70)
            ws.freeze_panes = "A2"
            if ws.title != "안내":
                ws.auto_filter.ref = ws.dimensions
            for col in ws.iter_cols():
                if "금액" in str(col[0].value):
                    for x in col[1:]:
                        x.number_format = "#,##0"
    return buf.getvalue()
